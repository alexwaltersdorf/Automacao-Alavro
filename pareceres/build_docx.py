#!/usr/bin/env python3
"""Gera versoes .docx (Word) dos documentos, com o mesmo layout dos PDFs.

Uso:
    python3 build_docx.py            # documentos da clinica (oficio e parecer colegiado)
    python3 build_docx.py --all      # todos os documentos listados em build_pdf.DOCS
    python3 build_docx.py arquivo.md # documento especifico

Requer: markdown, lxml, python-docx
"""
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Mm, Pt, RGBColor
from lxml import html as LH

from build_pdf import BASE, DOCS, LETTERHEAD, WORKHEAD, INTERNAL_BANNER, prepare, render

OUT = BASE / "docx"

BODY_FONT = "Georgia"
BODY_SIZE = Pt(10.5)
LINE = Pt(15.75)  # equivale ao line-height 1.5 do PDF
INK = RGBColor(0x14, 0x16, 0x1A)
MUTED = RGBColor(0x4B, 0x55, 0x63)

# ordem dos filhos de w:pPr no schema; usada para inserir w:pBdr na posicao certa
_PPR_AFTER_PBDR = (
    "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
    "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
    "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
    "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)


def border(par, edge="bottom", sz=6, color="14161A", space=4):
    """Aplica uma borda a um paragrafo (regua horizontal, cabecalho, assinatura)."""
    pPr = par._p.get_or_add_pPr()
    pbdr = pPr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        pPr.insert_element_before(pbdr, *_PPR_AFTER_PBDR)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color)
    pbdr.append(el)


def _numbering(doc):
    return doc.part.numbering_part.element


def abstract_of(doc, num_id):
    """Descobre o abstractNumId por tras do numId usado por um estilo de lista."""
    for n in _numbering(doc).findall(qn("w:num")):
        if n.get(qn("w:numId")) == str(num_id):
            return n.find(qn("w:abstractNumId")).get(qn("w:val"))
    return "0"


def fresh_num(doc, abstract_id, start=1):
    """Cria uma numeracao propria para a lista, para que ela reinicie do 1."""
    numbering = _numbering(doc)
    ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
    new_id = (max(ids) if ids else 0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_id))
    ab = OxmlElement("w:abstractNumId")
    ab.set(qn("w:val"), str(abstract_id))
    num.append(ab)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start_el = OxmlElement("w:startOverride")
    start_el.set(qn("w:val"), str(start))
    override.append(start_el)
    num.append(override)
    numbering.append(num)
    return new_id


def apply_num(par, num_id, ilvl=0):
    numPr = par._p.get_or_add_pPr().get_or_add_numPr()
    numPr.get_or_add_ilvl().val = ilvl
    numPr.get_or_add_numId().val = num_id


def style_num_id(doc, style_name):
    st = doc.styles[style_name].element
    pPr = st.find(qn("w:pPr"))
    numPr = pPr.find(qn("w:numPr")) if pPr is not None else None
    return numPr.find(qn("w:numId")).get(qn("w:val")) if numPr is not None else "0"


def shade(cell, color="EEF1F5"):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shd)


def setup(doc):
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Mm(297), Mm(210)
    sec.top_margin, sec.bottom_margin = Mm(22), Mm(20)
    sec.left_margin, sec.right_margin = Mm(22), Mm(20)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = BODY_SIZE
    normal.font.color.rgb = INK
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.element.rPr.rFonts.set(qn("w:cs"), BODY_FONT)
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing = LINE
    pf.space_after = Mm(3)
    pf.widow_control = True

    for name in ("List Bullet", "List Number", "List Bullet 2", "List Number 2"):
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        st.font.name = BODY_FONT
        st.font.size = BODY_SIZE
        st.paragraph_format.line_spacing = LINE
        st.paragraph_format.space_after = Mm(1.4)
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def para(doc, text=None, size=None, bold=False, italic=False, align=None,
         color=None, space_before=None, space_after=None, style=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if space_before is not None:
        p.paragraph_format.space_before = space_before
    if space_after is not None:
        p.paragraph_format.space_after = space_after
    if text:
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = size
        if color:
            r.font.color.rgb = color
    return p


def inline(par, node, bold=False, italic=False, fill=False, size=None, color=None):
    """Percorre o HTML inline (strong/em/br/span.fill) escrevendo runs no paragrafo."""
    def emit(text):
        if not text:
            return
        r = par.add_run(text)
        r.bold, r.italic = bold, italic
        if fill:
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        if size:
            r.font.size = size
        if color:
            r.font.color.rgb = color

    emit(node.text)
    for child in node:
        tag = child.tag
        cls = child.get("class", "")
        if tag == "br":
            par.add_run().add_break()
        elif tag in ("strong", "b"):
            inline(par, child, True, italic, fill, size, color)
        elif tag in ("em", "i"):
            inline(par, child, bold, True, fill, size, color)
        elif tag == "span" and "fill" in cls:
            inline(par, child, bold, italic, True, size, color)
        else:
            inline(par, child, bold, italic, fill, size, color)
        if child.tail:
            emit(child.tail)


def text_of(node):
    return "".join(node.itertext()).strip()


def add_list(doc, node, level=0):
    ordered = node.tag == "ol"
    base = "List Number" if ordered else "List Bullet"
    style = base if level == 0 else f"{base} {min(level + 1, 3)}"
    num_id = None
    if ordered:
        try:
            start = int(node.get("start", 1))
        except ValueError:
            start = 1
        num_id = fresh_num(doc, abstract_of(doc, style_num_id(doc, style)), start)
    for li in node.findall("li"):
        nested = [c for c in li if c.tag in ("ul", "ol")]
        p = doc.add_paragraph(style=style)
        if num_id is not None:
            apply_num(p, num_id)
        holder = LH.Element("span")
        holder.text = li.text
        for c in li:
            if c.tag in ("ul", "ol"):
                continue
            holder.append(c)
        inline(p, holder)
        for sub in nested:
            add_list(doc, sub, level + 1)


def add_table(doc, node):
    rows = node.xpath(".//tr")
    if not rows:
        return
    ncols = max(len(r.xpath("./th|./td")) for r in rows)
    table = doc.add_table(rows=0, cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r in rows:
        cells_src = r.xpath("./th|./td")
        row = table.add_row().cells
        for i, src in enumerate(cells_src):
            if i >= ncols:
                break
            cell = row[i]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = Pt(12)
            p.paragraph_format.space_after = Pt(0)
            inline(p, src, bold=(src.tag == "th"), size=Pt(9))
            if src.tag == "th":
                shade(cell)
    para(doc, space_after=Mm(2))


def add_block(doc, node):
    tag = node.tag
    cls = node.get("class", "")

    if tag == "div" and "letterhead" in cls:
        name = node.find_class("name")
        meta = node.find_class("meta")
        if name:
            para(doc, text_of(name[0]).upper(), size=Pt(12.5), bold=True,
                 align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Mm(1.5))
        if meta:
            p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Mm(6))
            p.paragraph_format.line_spacing = Pt(11)
            inline(p, meta[0], size=Pt(8.5), color=MUTED)
            border(p, sz=12)
        return

    if tag == "div" and "internal" in cls:
        p = para(doc, text_of(node).upper(), size=Pt(9), bold=True,
                 align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Mm(6),
                 color=RGBColor(0x7F, 0x1D, 0x1D))
        for edge in ("top", "bottom", "left", "right"):
            border(p, edge, sz=8, color="B91C1C", space=4)
        return

    if tag == "div" and "sig" in cls:
        line = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=Mm(9),
                    space_after=Mm(0))
        # regua de 78 mm centrada, como no PDF (area util de 168 mm)
        line.paragraph_format.left_indent = Mm(45)
        line.paragraph_format.right_indent = Mm(45)
        line.add_run(" ")
        border(line, sz=6, space=1)
        who = node.find_class("who")
        for w in who:
            para(doc, text_of(w), size=Pt(10), bold=True,
                 align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Mm(1))
        return

    if tag == "div" and "footer-note" in cls:
        p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=Mm(6))
        inline(p, node, italic=True, size=Pt(8), color=MUTED)
        border(p, "top", sz=4, color="9CA3AF")
        return

    if tag == "div" and "pagebreak" in cls:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        return

    if tag == "h1":
        p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=Mm(4),
                 style="Heading 1")
        p.style.font.name = BODY_FONT
        inline(p, node, bold=True, size=Pt(14), color=INK)
        for r in p.runs:
            r.text = r.text.upper()
        p.paragraph_format.keep_with_next = True
        return

    if tag in ("h2", "h3"):
        size = Pt(11.5) if tag == "h2" else Pt(10.5)
        p = para(doc, align=WD_ALIGN_PARAGRAPH.LEFT, style=f"Heading {tag[1]}",
                 space_before=Mm(8) if tag == "h2" else Mm(6),
                 space_after=Mm(3) if tag == "h2" else Mm(2))
        inline(p, node, bold=True, size=size,
               color=INK if tag == "h2" else RGBColor(0x1F, 0x29, 0x37))
        if tag == "h2":
            for r in p.runs:
                r.text = r.text.upper()
            border(p, sz=5, space=2)
        p.paragraph_format.keep_with_next = True
        return

    if tag in ("h4", "h5", "h6"):
        p = para(doc, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=Mm(4))
        inline(p, node, bold=True)
        p.paragraph_format.keep_with_next = True
        return

    if tag == "p":
        p = para(doc)
        # blocos com quebra manual (endereco, destinatario) nao sao justificados
        if node.findall("br"):
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        inline(p, node)
        return

    if tag in ("ul", "ol"):
        add_list(doc, node)
        return

    if tag == "table":
        add_table(doc, node)
        return

    if tag == "blockquote":
        for child in node:
            p = para(doc, space_after=Mm(1))
            p.paragraph_format.left_indent = Mm(6)
            inline(p, child, italic=True, color=RGBColor(0x37, 0x41, 0x51))
            border(p, "left", sz=12, color="9CA3AF", space=6)
        return

    if tag == "hr":
        p = para(doc, space_before=Mm(3), space_after=Mm(3))
        border(p, sz=4, color="9CA3AF")
        return

    if tag == "div":
        for child in node:
            add_block(doc, child)
        return


def build(src: Path, title: str, kind: str = "clinic") -> Path:
    md_text = prepare(src.read_text(encoding="utf-8"), keep_notes=kind != "clinic")
    body_html = render(md_text)
    head = LETTERHEAD if kind == "clinic" else WORKHEAD
    banner = INTERNAL_BANNER if kind == "internal" else ""
    root = LH.fragment_fromstring(f"<div>{head}{banner}{body_html}</div>")

    doc = Document()
    setup(doc)
    doc.core_properties.title = title
    doc.core_properties.language = "pt-BR"
    for node in root:
        add_block(doc, node)
    # o Word abre o documento com um paragrafo vazio no fim se o ultimo for tabela
    OUT.mkdir(exist_ok=True)
    path = OUT / (src.stem + ".docx")
    doc.save(path)
    return path


def main(argv):
    if "--all" in argv:
        selected = DOCS
    elif [a for a in argv if not a.startswith("-")]:
        wanted = {Path(a).name for a in argv if not a.startswith("-")}
        selected = [d for d in DOCS if d[0] in wanted]
        if not selected:
            sys.exit(f"nenhum documento corresponde a: {', '.join(sorted(wanted))}")
    else:
        selected = [d for d in DOCS if d[2] == "clinic"]

    for filename, title, kind in selected:
        src = BASE / filename
        if not src.exists():
            sys.exit(f"nao encontrado: {src}")
        print(f"gerado [{kind}]:", build(src, title, kind).name)


if __name__ == "__main__":
    main(sys.argv[1:])
